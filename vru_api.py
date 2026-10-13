"""Завантаження зареєстрованих законопроєктів з офіційного порталу Верховної Ради."""
from __future__ import annotations

import io
import logging
import re
import time
from datetime import datetime
from urllib.parse import urljoin, urlparse, parse_qs, urlencode, urlunparse

import requests
from bs4 import BeautifulSoup

try:
    from pypdf import PdfReader
except ImportError:  # PDF extraction is optional
    PdfReader = None

logger = logging.getLogger(__name__)
BASE_URL = "https://itd.rada.gov.ua"
PERIOD_URL = f"{BASE_URL}/billinfo/Bills/Period"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (compatible; VRULawMonitor/1.0; +https://github.com/skolesnicenko379-design/open_data_verhovnaradaua)",
    "Accept-Language": "uk-UA,uk;q=0.9,en;q=0.7",
}
SESSION = requests.Session()
SESSION.headers.update(HEADERS)
TIMEOUT = 25
MAX_PAGES = 8


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def get_soup(url: str) -> BeautifulSoup:
    response = SESSION.get(url, timeout=TIMEOUT)
    response.raise_for_status()
    response.encoding = response.apparent_encoding or "utf-8"
    return BeautifulSoup(response.text, "html.parser")


def _pagination_urls(soup: BeautifulSoup, current_url: str) -> list[str]:
    urls = []
    for anchor in soup.select("a[href]"):
        label = clean_text(anchor.get_text(" ", strip=True))
        href = urljoin(current_url, anchor["href"])
        if label.isdigit() or label in {"»", ">", "Наступна"}:
            if "/billinfo/Bills/Period" in href and href not in urls and href != current_url:
                urls.append(href)
    return urls


def _card_link(row) -> str | None:
    for anchor in row.select("a[href]"):
        href = urljoin(BASE_URL, anchor["href"])
        if re.search(r"/billinfo/Bills/Card/\d+", href, re.I):
            return href
    return None


def list_recent_bills() -> list[dict]:
    """Отримує законопроєкти з офіційної сторінки надходжень за поточний тиждень."""
    found: dict[str, dict] = {}
    queue = [PERIOD_URL]
    visited: set[str] = set()

    while queue and len(visited) < MAX_PAGES:
        url = queue.pop(0)
        if url in visited:
            continue
        visited.add(url)
        try:
            soup = get_soup(url)
        except requests.RequestException as exc:
            logger.error("Не вдалося завантажити перелік ВРУ %s: %s", url, exc)
            continue

        for row in soup.select("tr"):
            card_url = _card_link(row)
            if not card_url:
                continue
            row_text = clean_text(row.get_text(" ", strip=True))
            match = re.search(r"(?<!\d)(\d{4,6}(?:-\d+)?(?:/П\d*)?)(?!\d)", row_text)
            if not match:
                continue
            number = match.group(1)
            found[number] = {
                "number": number,
                "list_title": row_text,
                "card_url": card_url,
            }

        for next_url in _pagination_urls(soup, url):
            if next_url not in visited and next_url not in queue:
                queue.append(next_url)
        time.sleep(0.35)

    logger.info("У переліку ВРУ знайдено карток: %d", len(found))
    return list(found.values())


def _text_after_label(lines: list[str], label: str, stop_labels: tuple[str, ...] = ()) -> str:
    for i, line in enumerate(lines):
        if label.casefold() in line.casefold():
            parts = []
            inline = line.split(":", 1)
            if len(inline) == 2 and clean_text(inline[1]):
                parts.append(clean_text(inline[1]))
            for candidate in lines[i + 1:i + 5]:
                if any(stop.casefold() in candidate.casefold() for stop in stop_labels):
                    break
                if candidate:
                    parts.append(candidate)
            return clean_text(" ".join(parts))
    return ""


def _extract_linked_document(url: str) -> str:
    """Завантажує і витягує текст із PDF/DOCX або сторінки документа ВРУ."""
    def extract_file(response) -> str:
        content_type = response.headers.get("Content-Type", "").lower()
        payload = response.content
        if "pdf" in content_type or payload[:4] == b"%PDF":
            if PdfReader is None:
                return ""
            try:
                reader = PdfReader(io.BytesIO(payload))
                return clean_text("\n".join((page.extract_text() or "") for page in reader.pages[:25]))[:18000]
            except Exception as exc:
                logger.warning("Не вдалося розібрати PDF %s: %s", response.url, exc)
                return ""
        if "wordprocessingml" in content_type or payload[:2] == b"PK":
            try:
                from docx import Document
                document = Document(io.BytesIO(payload))
                paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
                for table in document.tables:
                    for row in table.rows:
                        paragraphs.append(" | ".join(cell.text for cell in row.cells))
                return clean_text("\n".join(paragraphs))[:18000]
            except Exception:
                return ""
        if "html" in content_type or b"<html" in payload[:500].lower():
            return ""
        return ""

    try:
        response = SESSION.get(url, timeout=TIMEOUT)
        response.raise_for_status()
        text = extract_file(response)
        if text:
            return text

        # Сторінка pubFile ВРУ показує кнопку «Завантажити»; сам PDF може бути
        # відданий окремим запитом форми, а не прямим посиланням.
        content_type = response.headers.get("Content-Type", "").lower()
        if "html" in content_type or b"<html" in response.content[:500].lower():
            soup = BeautifulSoup(response.content, "html.parser")
            for form in soup.find_all("form"):
                action = urljoin(response.url, form.get("action") or response.url)
                method = (form.get("method") or "get").lower()
                data = {}
                for field in form.select("input[name], select[name], textarea[name]"):
                    if field.name == "input" and field.get("type", "").lower() in {"submit", "button"}:
                        continue
                    data[field["name"]] = field.get("value", "")
                for button in form.select("button[formaction], input[formaction]"):
                    action = urljoin(response.url, button.get("formaction"))
                    if button.get("name"):
                        data[button["name"]] = button.get("value", "")
                try:
                    downloaded = SESSION.request(method, action, data=data if method == "post" else None,
                                                 params=data if method != "post" else None, timeout=TIMEOUT)
                    downloaded.raise_for_status()
                    text = extract_file(downloaded)
                    if text:
                        return text
                except Exception:
                    continue
            for anchor in soup.select("a[href]"):
                label = clean_text(anchor.get_text(" ", strip=True)).casefold()
                href = urljoin(response.url, anchor["href"])
                if "завантажити" in label or "download" in label or "download" in href.casefold():
                    try:
                        downloaded = SESSION.get(href, timeout=TIMEOUT)
                        downloaded.raise_for_status()
                        text = extract_file(downloaded)
                        if text:
                            return text
                    except Exception:
                        continue
        return ""
    except Exception as exc:
        logger.warning("Не вдалося прочитати супровідний документ %s: %s", url, exc)
        return ""

def get_bill_details(item: dict) -> dict | None:
    url = item["card_url"]
    soup = get_soup(url)
    lines = [clean_text(line) for line in soup.get_text("\n", strip=True).splitlines()]
    lines = [line for line in lines if line]
    full_text = "\n".join(lines)

    title = ""
    for line in lines:
        if line.startswith(("Проєкт Закону", "Проект Закону")):
            title = line
            break
    if not title:
        title = clean_text(item.get("list_title", ""))
    # Не публікуємо постанови, кадрові рішення та інші документи, що не є законопроєктами.
    if not re.match(r"^(Проєкт|Проект) Закону\b", title, re.I):
        return None

    number = item["number"]
    registration_date = ""
    number_date = re.search(
        r"Номер, дата реєстрації:\s*(?:\n\s*)?([\d-]+)\s+від\s+(\d{2}\.\d{2}\.\d{4})",
        full_text,
        re.I,
    )
    if number_date:
        number = number_date.group(1)
        registration_date = number_date.group(2)
    else:
        fallback = re.search(r"(?<!\d)(\d{4,6}(?:-\d+)?)(?:\s+від\s+)(\d{2}\.\d{2}\.\d{4})", full_text)
        if fallback:
            number, registration_date = fallback.groups()

    authors = []
    in_authors = False
    for anchor in soup.select("a[href]"):
        href = anchor.get("href", "")
        anchor_name = clean_text(anchor.get_text(" ", strip=True))
        if "Ініціатор(и) законопроєкту" in anchor_name or "Ініціатор(и) законопроєкту" in href:
            in_authors = True
        if re.search(r"/struct/uk/Structure/MPs\?userId=", href, re.I):
            name = re.sub(r"\s*\(IX скликання\)\s*", "", anchor_name).strip(" ,")
            if name and name not in authors:
                authors.append(name)
        # Картка має профільні посилання авторів; збір усіх депутатських посилань тут безпечний,
        # оскільки інші секції картки зазвичай містять комітети як звичайний текст.
    author_profiles = []
    for anchor in soup.select('a[href*="/struct/uk/Structure/MPs?userId="]'):
        name = re.sub(r"\s*\(IX скликання\)\s*", "", clean_text(anchor.get_text(" ", strip=True))).strip(" ,")
        href = urljoin(BASE_URL, anchor["href"])
        if name and all(p["name"] != name for p in author_profiles):
            author_profiles.append({"name": name, "url": href})
    if author_profiles:
        authors = [p["name"] for p in author_profiles]

    sponsor_text = _text_after_label(
        lines, "Суб'єкт права законодавчої ініціативи:",
        ("Ініціатор(и) законопроєкту:", "Головний комітет:")
    )
    explanatory_url = ""
    bill_text_url = ""
    for anchor in soup.select("a[href]"):
        label = clean_text(anchor.get_text(" ", strip=True)).casefold()
        href = urljoin(url, anchor["href"])
        if "пояснювальна записка" in label and not explanatory_url:
            explanatory_url = href
        elif ("проєкт закону" in label or "проект закону" in label) and not bill_text_url:
            bill_text_url = href

    explanatory_text = _extract_linked_document(explanatory_url) if explanatory_url else ""
    bill_text = _extract_linked_document(bill_text_url) if bill_text_url else ""
    source_text = (explanatory_text + "\n" + bill_text).strip()
    if not source_text:
        # Залишаємо структуровану частину картки як контекст, не видаючи її за повний текст законопроєкту.
        source_text = _text_after_label(lines, "Скорочений текст:", ("Формалізований текст:", "Поділитись:"))

    return {
        "number": number,
        "date": registration_date,
        "title": title,
        "card_url": url,
        "authors": authors,
        "author_profiles": author_profiles,
        "sponsor": sponsor_text,
        "source_text": source_text[:18000],
        "source_text_type": "супровідні документи" if explanatory_text or bill_text else "картка законопроєкту",
    }


def get_author_affiliation(profile_url: str) -> str:
    """Повертає партію/фракцію з офіційного профілю депутата."""
    try:
        soup = get_soup(profile_url)
        lines = [clean_text(line) for line in soup.get_text("\n", strip=True).splitlines()]
        lines = [line for line in lines if line]
        party = _text_after_label(lines, "ПАРТІЯ:")
        faction = _text_after_label(lines, "ФРАКЦІЯ:")
        if faction:
            faction = re.sub(r"^депутатської фракції\s+", "", faction, flags=re.I)
            return faction
        return party or "фракцію/партію не вказано в профілі ВРУ"
    except Exception as exc:
        logger.warning("Не вдалося завантажити профіль депутата %s: %s", profile_url, exc)
        return "фракцію/партію не вдалося підтвердити"
